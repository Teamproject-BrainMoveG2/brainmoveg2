from fastapi.responses import FileResponse
import socketio
import config
from typing_extensions import Annotated
from fastapi import APIRouter, Depends, HTTPException
from services.cone_service import ConeService
from services.game_service import GameService
from dependencies import get_cone_service, get_settings, get_sio
from repositories.mode_repository import ModeRepository
from repositories.tutorial_repository import TutorialRepository
from repositories.gamesession_repository import GameSessionRepository
from models.models import Cone, ConeStatusDTO, ExportDateDTO
import logging
import tempfile
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

logger = logging.getLogger(__name__)

router = APIRouter(
    prefix="/data",
    tags=["data"],
    dependencies=[Depends(get_sio)],
    responses={404: {"description": "Not found"}},
)

@router.get("/today")
async def get_today_data(settings: Annotated[config.Settings, Depends(get_settings)]):
    try:
        data = GameSessionRepository.get_sessions_today(settings)
    except Exception as e:
        logger.error(f"Error retrieving today's data: {e}")
        raise HTTPException(status_code=400, detail=str(e))
    return {"data": data}

@router.get("/week")
async def get_week_data(settings: Annotated[config.Settings, Depends(get_settings)]):
    try:
        data = GameSessionRepository.get_sessions_thisweek(settings)
    except Exception as e:
        logger.error(f"Error retrieving this week's data: {e}")
        raise HTTPException(status_code=400, detail=str(e))
    return {"data": data}

@router.get("/month")
async def get_month_data(settings: Annotated[config.Settings, Depends(get_settings)]):
    try:
        data = GameSessionRepository.get_sessions_thismonth(settings)
    except Exception as e:
        logger.error(f"Error retrieving this month's data: {e}")
        raise HTTPException(status_code=400, detail=str(e))
    return {"data": data}

@router.get("/alltime")
async def get_alltime_data(settings: Annotated[config.Settings, Depends(get_settings)]):
    try:
        data = GameSessionRepository.get_sessions_alltime(settings)
    except Exception as e:
        logger.error(f"Error retrieving all-time data: {e}")
        raise HTTPException(status_code=400, detail=str(e))
    return {"data": data}

@router.get("/excel")
async def get_excel_data(settings: Annotated[config.Settings, Depends(get_settings)]):
    range_val = 'alltime'
    if range_val == 'today':
        data = GameSessionRepository.get_sessions_today(settings)
    elif range_val == 'week':
        data = GameSessionRepository.get_sessions_thisweek(settings)
    elif range_val == 'month':
        data = GameSessionRepository.get_sessions_thismonth(settings)
    else:
        data = GameSessionRepository.get_sessions_alltime(settings)

    wb = Workbook()
    ws = wb.active
    ws.title = "BrainMove Data"

    # 1. Define Headers and Keys
    headers = ["ID", "Gebruikersnaam", "Moeilijkheid", "Score", "Gem. Reactietijd (ms)", "Accuraatheid (%)"]
    keys = ["spelsessie_id", "username", "naam", "score", "avg_reactietijd_ms", "accuracy_percent"]

    # 2. Add Header Row
    ws.append(headers)

    # 3. Styling Definitions
    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    for col_num, cell in enumerate(ws[1], 1):
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    # 5. add Data Rows
    for row_data in data:
        row = [row_data.get(key) for key in keys]
        ws.append(row)

    # 6. Apply styling to data rows and Auto-size columns
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.alignment = center_align
            cell.border = thin_border
            # Optional: Add specific number formats
            if cell.column == 6: # Accuracy column
                 cell.number_format = '0.0"%"' # Visually display matches as %

    # Auto-adjust column widths
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter # Get the column name
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column].width = adjusted_width
    # Temp file logic
    with tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx') as tmp:
        wb.save(tmp.name)
        file_path = tmp.name

    return FileResponse(
        path=file_path,
        filename='brainmove_export.xlsx',
        media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': 'attachment; filename="brainmove_export.xlsx"'}
    )
    

