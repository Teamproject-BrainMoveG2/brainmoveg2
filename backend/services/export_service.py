import tempfile
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

class ExportService:
    def export_data_to_excel(self, data):
        wb = Workbook()
        ws = wb.active
        ws.title = "BrainMove Data"

        headers = ["Datum", "Gebruikersnaam", "Moeilijkheid", "Score", "Gem. Reactietijd (ms)", "Accuraatheid (%)"]
        keys = ["datum", "username", "naam", "score", "avg_reactietijd_ms", "accuracy_percent"]

        ws.append(headers)

        header_font = Font(bold=True, color="FFFFFF", size=12)
        header_fill = PatternFill(start_color="10A8C9", end_color="10A8C9", fill_type="solid")
        align = Alignment(horizontal="left", vertical="center")
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

        for col_num, cell in enumerate(ws[1], 1):
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = align
            cell.border = thin_border

        for row_data in data:
            row = [row_data.get(key) for key in keys]
            ws.append(row)

        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.alignment = align
                cell.border = thin_border
                if cell.column == 6:
                    cell.number_format = '0.0"%"'

        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    val_len = len(str(cell.value))
                    if cell.row == 1:
                        val_len = int(val_len * 1.3) 

                    if val_len > max_length:
                        max_length = val_len
                except:
                    pass
            adjusted_width = (max_length + 2)
            ws.column_dimensions[column].width = adjusted_width

        with tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx') as tmp:
            wb.save(tmp.name)
            file_path = tmp.name

        return file_path